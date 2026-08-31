"""Resume Intelligence service — safe extraction + structured parsing.

FEATURE 1 (part 1): reads PDF/DOCX/TXT uploads and derives a
:class:`ResumeProfile` (contact details, sections, skills, education,
experience and projects). Parsing is deterministic rule-based logic —
no LLM required — so results are reproducible and private.

Binary parsers (pypdf / python-docx) are imported lazily so the module
imports cleanly even before `pip install pypdf python-docx`.
"""
from __future__ import annotations

import io
import re

from dashboard.models.resume import ResumeProfile
from dashboard.utils.constants import RESUME_SECTION_HEADINGS
from dashboard.utils.text_processing import (
    detect_education_level,
    extract_years_experience,
    normalize_whitespace,
)
from dashboard.utils.validation import ResumeParseError, validate_upload

_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?:\+?\d[\d\s\-()]{7,}\d)")
_HEADING_RE = re.compile(
    r"^(?:{0})\s*:?\s*$".format("|".join(h for h in RESUME_SECTION_HEADINGS)),
    re.IGNORECASE,
)
# Fallback: short standalone lines like "EXPERIENCE:" treated as headings
_SECTION_KEY_RE = re.compile(
    r"({0})\s*:?".format("|".join(h for h in RESUME_SECTION_HEADINGS)),
    re.IGNORECASE,
)
_CERT_KEYWORDS = re.compile(
    r"certifi|certified|aws|azure|gcp|coursera|udemy|nptel|"
    r"google data analytics|hubspot|tableau desktop",
    re.IGNORECASE,
)
_DEGREE_LINE_RE = re.compile(
    r"b\.?tech|b\.?e\b|bachelor|master|m\.?tech|mba|m\.?c\.?a|b\.?c\.?a|"
    r"b\.?s\.?c|m\.?s\.?c|b\.?com|m\.?com|ph\.?d|diploma",
    re.IGNORECASE,
)


def extract_text(filename: str, data: bytes) -> str:
    """Extract raw text from a PDF/DOCX/TXT upload (validates first)."""
    ext = validate_upload(filename, data)

    if ext == ".txt":
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1", errors="ignore")
        return normalize_whitespace(text)

    if ext == ".pdf":
        return _extract_pdf(data)
    return _extract_docx(data)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader  # lazy import (see module docstring)
    except ImportError as exc:  # pragma: no cover - dependency guarded
        raise ResumeParseError(
            "PDF support is not installed on the server. Run: pip install pypdf"
        ) from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - library raises many types
        raise ResumeParseError(
            "This PDF could not be read. It may be corrupted or password-protected. "
            "Try re-exporting it as a standard PDF or upload a DOCX/TXT version."
        ) from exc
    text = normalize_whitespace("\n".join(pages))
    if len(text.strip()) < 40:
        raise ResumeParseError(
            "No selectable text was found in this PDF (it may be a scan). "
            "Please upload a text-based PDF, DOCX or TXT resume."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx, lazy import
    except ImportError as exc:  # pragma: no cover - dependency guarded
        raise ResumeParseError(
            "DOCX support is not installed on the server. Run: pip install python-docx"
        ) from exc
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except Exception as exc:  # noqa: BLE001
        raise ResumeParseError(
            "This DOCX file could not be read. Please re-save it in Word format "
            "or upload a PDF/TXT version."
        ) from exc
    return normalize_whitespace("\n".join(parts))


def _split_sections(text: str) -> dict[str, str]:
    """Split resume text into named sections using standard headings."""
    sections: dict[str, str] = {}
    current = "header"
    buffer: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if _HEADING_RE.match(stripped) or (
            0 < len(stripped) < 40 and _SECTION_KEY_RE.fullmatch(stripped.rstrip(":"))
        ):
            key = stripped.rstrip(":").strip().lower()
            if buffer:
                sections[current] = "\n".join(buffer).strip()
            current = key or "other"
            buffer = []
        else:
            buffer.append(line)
    if buffer:
        sections[current] = "\n".join(buffer).strip()
    return sections


def _detect_name(lines: list[str]) -> str:
    """Heuristic: first short line of 2-4 capitalised words without '@'."""
    for line in lines[:4]:
        clean = line.strip().strip("|,")
        if not clean or "@" in clean or any(ch.isdigit() for ch in clean):
            continue
        tokens = [t for t in re.split(r"\s+", clean) if t]
        if 2 <= len(tokens) <= 4 and all(t[0].isupper() and t.isalpha() for t in tokens):
            return clean
    return ""


def parse_resume(text: str, filename: str = "") -> ResumeProfile:
    """Parse raw resume text into a structured ResumeProfile."""
    if not text or len(text.strip()) < 60:
        raise ResumeParseError(
            "The resume text is too short to analyse. Please upload a complete resume."
        )

    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    sections = _split_sections(text)

    email_m = _EMAIL_RE.search(text)
    compact = text.replace(" ", "")
    phone_m = _PHONE_RE.search(compact) or _PHONE_RE.search(text)

    from src.skill_extractor import extract_skills_from_text  # existing engine

    skills_blob = " ".join(
        sections.get(k, "") for k in ("skills", "technical skills")
    )
    skills = extract_skills_from_text(skills_blob) or extract_skills_from_text(text)

    edu_source = sections.get("education", "") or "\n".join(lines)
    education_lines = [
        ln.strip() for ln in edu_source.splitlines() if _DEGREE_LINE_RE.search(ln)
    ][:4]

    projects_section = sections.get("projects", "")
    projects = [
        re.sub(r"^[•\-\*\\d.)\s]+", "", ln.strip())[:80]
        for ln in projects_section.splitlines()
        if ln.strip() and len(ln.strip()) > 8
    ][:8]

    cert_source = (
        sections.get("certifications", "") or sections.get("certificates", "")
    )
    certifications = [
        ln.strip()[:90] for ln in cert_source.splitlines()
        if ln.strip() and len(ln.strip()) > 6
    ][:6]
    if not certifications:
        certifications = [
            ln.strip()[:90] for ln in lines
            if _CERT_KEYWORDS.search(ln) and not _EMAIL_RE.search(ln)
        ][:4]

    experience_years = extract_years_experience(
        sections.get("experience", "") or sections.get("work experience", "") or text
    )

    profile = ResumeProfile(
        filename=filename,
        text=text,
        name=_detect_name(lines),
        email=email_m.group(0) if email_m else "",
        phone=phone_m.group(0) if phone_m else "",
        skills=skills,
        education=[ln for ln in education_lines if ln],
        education_level=detect_education_level(text),
        experience_years=experience_years,
        projects=projects,
        certifications=certifications,
        sections_found=sorted(sections.keys()),
        word_count=len(text.split()),
    )

    if not profile.email:
        profile.warnings.append("No email address detected — recruiters cannot reach you.")
    if not profile.phone:
        profile.warnings.append("No phone number detected.")
    if not profile.skills:
        profile.warnings.append(
            "No recognisable skills found — list tools/technologies explicitly."
        )
    if profile.word_count < 120:
        profile.warnings.append("Resume looks very short — aim for 300+ words.")
    return profile


def analyze_resume(filename: str, data: bytes) -> ResumeProfile:
    """Extract + parse an uploaded resume (uncached; pages add caching)."""
    return parse_resume(extract_text(filename, data), filename=filename)
